"""
HYBRID AGENT - Nexus tools
==========================

The Agent Nexus abilities (JarvisAgent/nexus_actions), rebuilt for the web
console's agent mode. The model picks a tool; serve.py runs it.

Every tool has a risk tier:
    auto -> read-only, runs as soon as the model asks for it
    ask  -> changes something on this PC; the browser shows Approve / Deny
            and only an approved call reaches /api/agent/tool

Desktop tools only exist on a local server. On a cloud deploy (RENDER set)
they would let any visitor drive the server itself, so `available()` is
empty there and the agent can only talk.

Heavy or Windows-only libraries are imported inside each tool, so a missing
package breaks that one tool, not the server.
"""

from __future__ import annotations

import base64
import io
import json
import os
import shlex
import shutil
import subprocess
import sys
import time
import urllib.request
import webbrowser
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

HERE = Path(__file__).resolve().parent
OUT_DIR = Path.home() / "Documents" / "Hybrid Agent"
MEMORY_FILE = HERE / "nexus_memory.json"
IGNORE_DIRS = {".git", "__pycache__", ".venv", "venv", "node_modules", ".idea", ".vscode", "dist", "build"}
MAX_RESULT = 6000


def local_mode() -> bool:
    return not os.environ.get("RENDER")


@dataclass
class Tool:
    name: str
    summary: str
    params: dict = field(default_factory=dict)      # arg name -> description
    risk: str = "auto"                               # "auto" | "ask"
    fn: Callable[..., str] | None = None


TOOLS: dict[str, Tool] = {}


def tool(name: str, summary: str, params: dict | None = None, risk: str = "auto"):
    def wrap(fn):
        TOOLS[name] = Tool(name, summary, params or {}, risk, fn)
        return fn
    return wrap


# ----------------------------------------------------------------- helpers
def _path(p: str, base: Path = HERE) -> Path:
    path = Path(os.path.expandvars(os.path.expanduser(str(p or ".").strip().strip('"'))))
    return (path if path.is_absolute() else base / path).resolve()


def _safe_name(text: str, ext: str) -> str:
    stem = "".join(c if c.isalnum() or c in " _-" else "" for c in text).strip().replace(" ", "_")
    return (stem[:60] or "untitled") + ext


def _unique(path: Path) -> Path:
    n, out = 2, path
    while out.exists():
        out = path.with_name(f"{path.stem}_{n}{path.suffix}")
        n += 1
    return out


def _open_file(path: Path) -> None:
    try:
        if os.name == "nt":
            os.startfile(str(path))                  # noqa: S606 - our own output file
    except OSError:
        pass


def _int(v, default: int, lo: int, hi: int) -> int:
    try:
        return max(lo, min(hi, int(v)))
    except (TypeError, ValueError):
        return default


# ----------------------------------------------------------------- memory
def _memory() -> dict:
    try:
        return json.loads(MEMORY_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"rules": [], "history": []}


def _save_memory(mem: dict) -> None:
    MEMORY_FILE.write_text(json.dumps(mem, indent=2), encoding="utf-8")


def memory_context() -> str:
    rules = _memory().get("rules", [])
    if not rules:
        return ""
    return "Things the user asked you to remember:\n" + "\n".join(f"- {r}" for r in rules)


@tool("remember", "Save a preference or fact about the user for future conversations.",
      {"rule": "the thing to remember, one sentence"})
def remember(rule: str = "") -> str:
    rule = str(rule).strip()
    if not rule:
        return "Nothing to remember."
    mem = _memory()
    rules = mem.setdefault("rules", [])
    if rule not in rules:
        rules.append(rule)
        _save_memory(mem)
    return f"Remembered: {rule}"


@tool("forget", "Remove a remembered item by its number (1 = first).", {"number": "which item"}, risk="ask")
def forget(number=0) -> str:
    mem = _memory()
    rules = mem.get("rules", [])
    i = _int(number, 0, 0, 10_000) - 1
    if not 0 <= i < len(rules):
        return f"There is no item {number}. Remembered items: {len(rules)}."
    gone = rules.pop(i)
    _save_memory(mem)
    return f"Forgot: {gone}"


# ----------------------------------------------------------------- system
@tool("system_vitals", "CPU, RAM, disk and battery right now.")
def system_vitals() -> str:
    import psutil
    mem = psutil.virtual_memory()
    disk = psutil.disk_usage(str(Path.home().anchor or "/"))
    bat = psutil.sensors_battery()
    battery = f"{bat.percent:.0f}% ({'plugged in' if bat.power_plugged else 'on battery'})" if bat else "n/a"
    return (f"CPU {psutil.cpu_percent(interval=0.3)}%\n"
            f"RAM {mem.percent}% ({mem.used // 2**20} MB of {mem.total // 2**20} MB)\n"
            f"Disk {disk.percent}% used, {disk.free // 2**30} GB free\n"
            f"Battery {battery}")


@tool("memory_hogs", "The processes using the most RAM.", {"top_n": "how many (default 5)"})
def memory_hogs(top_n=5) -> str:
    import psutil
    rows = []
    for p in psutil.process_iter(["pid", "name", "memory_info"]):
        try:
            rows.append((p.info["memory_info"].rss / 2**20, p.info["name"], p.info["pid"]))
        except (psutil.NoSuchProcess, psutil.AccessDenied, AttributeError):
            continue
    rows.sort(reverse=True)
    return "\n".join(f"{name} (PID {pid}): {mb:.0f} MB" for mb, name, pid in rows[:_int(top_n, 5, 1, 25)])


@tool("kill_process", "Force-close an app by process name (e.g. 'notepad') or PID.",
      {"target": "process name or PID"}, risk="ask")
def kill_process(target: str = "") -> str:
    import psutil
    target = str(target).strip().lower()
    if not target:
        return "Say which process to close."
    names = {target, target if target.endswith(".exe") or target.isdigit() else target + ".exe"}
    protected = {"explorer.exe", "winlogon.exe", "csrss.exe", "lsass.exe", "services.exe", "system", "svchost.exe"}
    killed = []
    for p in psutil.process_iter(["pid", "name"]):
        try:
            name = (p.info["name"] or "").lower()
            if (str(p.info["pid"]) == target or name in names) and name not in protected and p.pid != os.getpid():
                p.kill()
                killed.append(f"{p.info['name']} (PID {p.info['pid']})")
        except (psutil.NoSuchProcess, psutil.AccessDenied):
            continue
    return f"Closed: {', '.join(killed)}" if killed else f"No running process matched '{target}'."


# ----------------------------------------------------------------- desktop organizer
CATEGORIES = {
    "Documents": {".pdf", ".docx", ".doc", ".txt", ".xlsx", ".xls", ".pptx", ".csv", ".odt", ".rtf", ".md"},
    "Images & Media": {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg", ".webp", ".mp4", ".mp3", ".wav", ".mkv", ".mov"},
    "Installers": {".exe", ".msi", ".iso"},
    "Archives": {".zip", ".rar", ".7z", ".tar", ".gz"},
    "Code & Scripts": {".py", ".js", ".html", ".css", ".json", ".ts", ".cpp", ".c", ".java", ".sh", ".bat", ".ps1"},
}


def _desktop() -> Path:
    return Path.home() / "Desktop"


@tool("organize_desktop", "Sort loose files on the Desktop into category folders. Can be undone.", risk="ask")
def organize_desktop() -> str:
    desk = _desktop()
    manifest = desk / ".hybrid_organizer_undo.json"
    moves = []
    for item in desk.iterdir():
        if item.is_dir() or item.name.startswith(".") or item.suffix.lower() in {".lnk", ".url", ".ini"}:
            continue
        cat = next((c for c, exts in CATEGORIES.items() if item.suffix.lower() in exts), "Other")
        dest = _unique(desk / cat / item.name)
        dest.parent.mkdir(exist_ok=True)
        shutil.move(str(item), str(dest))
        moves.append({"from": str(item), "to": str(dest)})
    if not moves:
        return "The Desktop has no loose files to sort."
    previous = json.loads(manifest.read_text(encoding="utf-8")) if manifest.exists() else []
    manifest.write_text(json.dumps(previous + moves, indent=2), encoding="utf-8")
    return f"Moved {len(moves)} file(s) into category folders. Say 'undo the desktop sort' to put them back."


@tool("undo_organize_desktop", "Put every file the desktop sort moved back where it was.", risk="ask")
def undo_organize_desktop() -> str:
    desk = _desktop()
    manifest = desk / ".hybrid_organizer_undo.json"
    if not manifest.exists():
        return "There is no desktop sort to undo."
    moves = json.loads(manifest.read_text(encoding="utf-8"))
    restored, folders = 0, set()
    for m in reversed(moves):
        src, dst = Path(m["to"]), Path(m["from"])
        folders.add(src.parent)
        if src.exists() and not dst.exists():
            shutil.move(str(src), str(dst))
            restored += 1
    for f in folders:
        try:
            f.rmdir()                                   # only succeeds when empty
        except OSError:
            pass
    manifest.unlink()
    return f"Put {restored} file(s) back on the Desktop."


# ----------------------------------------------------------------- documents
@tool("create_presentation", "Build a PowerPoint (.pptx) deck and open it.",
      {"title": "deck title",
       "slides": 'JSON list like [{"title": "...", "bullets": ["...", "..."]}], 3-8 slides'},
      risk="ask")
def create_presentation(title: str = "Presentation", slides="[]") -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    data = json.loads(slides) if isinstance(slides, str) else slides
    if not isinstance(data, list) or not data:
        return "No slides given. Pass a list of {title, bullets}."
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def slide(heading: str, bullets: list, cover: bool = False):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(16, 19, 14)
        box = s.shapes.add_textbox(Inches(0.9), Inches(2.6 if cover else 0.7), Inches(11.5), Inches(1.4))
        box.text_frame.word_wrap = True
        p = box.text_frame.paragraphs[0]
        p.text = str(heading)
        p.font.size, p.font.bold = Pt(48 if cover else 34), True
        p.font.color.rgb = RGBColor(200, 243, 107)
        if bullets:
            body = s.shapes.add_textbox(Inches(0.9), Inches(4.0 if cover else 2.1), Inches(11.5), Inches(4.6))
            tf = body.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets):
                bp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                bp.text = ("" if cover else "•  ") + str(b)
                bp.font.size = Pt(22 if cover else 20)
                bp.font.color.rgb = RGBColor(236, 236, 228)
                bp.space_after = Pt(12)

    slide(title, [datetime.now().strftime("%d %B %Y")], cover=True)
    for i, s in enumerate(data):
        s = s if isinstance(s, dict) else {"title": str(s)}
        slide(s.get("title") or f"Slide {i + 1}", s.get("bullets") or [])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = _unique(OUT_DIR / _safe_name(title, ".pptx"))
    prs.save(str(out))
    _open_file(out)
    return f"Saved {len(data) + 1} slides to {out}\n[file:{out.name}]"


@tool("create_word_document", "Write a Word (.docx) document and open it.",
      {"title": "document title",
       "content": "the body; blank lines split paragraphs, lines starting '- ' are bullets, '# ' are headings"},
      risk="ask")
def create_word_document(title: str = "Document", content: str = "") -> str:
    import docx
    doc = docx.Document()
    doc.add_heading(str(title), level=0)
    for block in str(content).split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("#"):
            doc.add_heading(block.lstrip("# ").strip(), level=1)
        elif block.startswith(("- ", "* ")):
            for line in block.splitlines():
                doc.add_paragraph(line.strip().lstrip("-* "), style="List Bullet")
        else:
            doc.add_paragraph(block)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = _unique(OUT_DIR / _safe_name(title, ".docx"))
    doc.save(str(out))
    _open_file(out)
    return f"Saved the document to {out}\n[file:{out.name}]"


# ----------------------------------------------------------------- code + repos
@tool("explore_folder", "Show the file tree of a folder (a code repo, a project, Documents...).",
      {"path": "folder path; '.' is the Hybrid Agent folder, '~' is the home folder", "depth": "levels (default 3)"})
def explore_folder(path: str = ".", depth=3) -> str:
    root = _path(path)
    if not root.is_dir():
        return f"Not a folder: {root}"
    lines, limit = [f"{root}"], _int(depth, 3, 1, 6)

    def walk(d: Path, prefix: str, level: int):
        if level >= limit or len(lines) > 150:
            return
        try:
            entries = sorted((e for e in d.iterdir() if e.name not in IGNORE_DIRS),
                             key=lambda e: (not e.is_dir(), e.name.lower()))
        except OSError:
            return
        for i, e in enumerate(entries):
            last = i == len(entries) - 1
            lines.append(f"{prefix}{'└── ' if last else '├── '}{e.name}{'/' if e.is_dir() else ''}")
            if e.is_dir():
                walk(e, prefix + ("    " if last else "│   "), level + 1)

    walk(root, "", 0)
    return "\n".join(lines[:150])


@tool("search_files", "Find text inside files in a folder.",
      {"path": "folder to search", "text": "what to look for", "extensions": "optional, e.g. 'py,js'"})
def search_files(path: str = ".", text: str = "", extensions: str = "") -> str:
    root, needle = _path(path), str(text).lower()
    if not needle:
        return "Say what text to look for."
    exts = {e.strip().lower().lstrip(".") for e in str(extensions).split(",") if e.strip()}
    hits = []
    for p in root.rglob("*"):
        if len(hits) >= 40:
            break
        if not p.is_file() or any(part in IGNORE_DIRS for part in p.parts):
            continue
        if exts and p.suffix.lstrip(".").lower() not in exts:
            continue
        if p.stat().st_size > 2_000_000:
            continue
        try:
            with open(p, encoding="utf-8", errors="ignore") as f:
                for n, line in enumerate(f, 1):
                    if needle in line.lower():
                        hits.append(f"{p.relative_to(root)}:{n}: {line.strip()[:120]}")
                        if len(hits) >= 40:
                            break
        except OSError:
            continue
    return "\n".join(hits) if hits else f"No match for '{text}' in {root}."


@tool("read_file", "Read a text file.", {"path": "file path", "max_lines": "default 200"})
def read_file(path: str = "", max_lines=200) -> str:
    p = _path(path)
    if not p.is_file():
        return f"File not found: {p}"
    limit, out = _int(max_lines, 200, 1, 1000), []
    with open(p, encoding="utf-8", errors="replace") as f:
        for i, line in enumerate(f):
            if i >= limit:
                out.append(f"... (stopped at {limit} lines)")
                break
            out.append(line.rstrip("\n"))
    return f"{p}\n" + "\n".join(out)


@tool("save_code_file", "Write code to a file (creates folders). Python is syntax-checked first.",
      {"path": "where to save; relative paths go in Documents/Hybrid Agent", "code": "the full file contents"},
      risk="ask")
def save_code_file(path: str = "", code: str = "") -> str:
    if not str(path).strip():
        return "Say where to save the file."
    p = _path(path, OUT_DIR)
    code = str(code)
    if code.lstrip().startswith("```"):
        code = code.split("\n", 1)[1] if "\n" in code else ""
        code = code.rsplit("```", 1)[0]
    if p.suffix == ".py":
        import ast
        try:
            ast.parse(code)
        except SyntaxError as e:
            return f"Not saved: SyntaxError on line {e.lineno}: {e.msg}"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(code, encoding="utf-8")
    return f"Saved {len(code.splitlines())} lines to {p}"


@tool("run_script", "Run a script file (.py, .js, .bat, .ps1) and return its output (60s limit).",
      {"path": "script path"}, risk="ask")
def run_script(path: str = "") -> str:
    p = _path(path, OUT_DIR)
    if not p.is_file():
        return f"File not found: {p}"
    runners = {".py": [sys.executable, str(p)], ".js": ["node", str(p)], ".bat": ["cmd", "/c", str(p)],
               ".cmd": ["cmd", "/c", str(p)],
               ".ps1": ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(p)]}
    cmd = runners.get(p.suffix.lower())
    if not cmd:
        return f"Cannot run {p.suffix} files."
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=60, cwd=str(p.parent))
    except subprocess.TimeoutExpired:
        return "Stopped: the script ran for more than 60 seconds."
    return f"exit code {r.returncode}\n{r.stdout.strip()}\n{r.stderr.strip()}".strip()


GIT_READ = {"status", "log", "diff", "show", "branch"}
GIT_WRITE = {"add", "commit", "stash", "checkout", "switch", "restore", "pull", "fetch", "init"}


def _git(command: str, path: str) -> str:
    args = shlex.split(str(command).strip().removeprefix("git ").strip(), posix=os.name != "nt")
    if not args:
        return "Say which git command to run."
    if any("output" in a or "ext-diff" in a or "textconv" in a or a.startswith("--exec") for a in args):
        return "That git option is not allowed here."
    cwd = _path(path)
    try:
        r = subprocess.run(["git", "--no-pager", *args], capture_output=True, text=True, timeout=30, cwd=str(cwd))
    except FileNotFoundError:
        return "git is not installed."
    return (r.stdout.strip() or r.stderr.strip() or "done.")


@tool("git_read", "Look at a git repo: status, log, diff, show or branch.",
      {"command": "e.g. 'status' or 'log --oneline -10'", "path": "repo folder"})
def git_read(command: str = "status", path: str = ".") -> str:
    sub = (shlex.split(str(command).removeprefix("git ").strip() or "status") or ["status"])[0]
    if sub not in GIT_READ:
        return f"git_read only runs {', '.join(sorted(GIT_READ))}. Use git_change for '{sub}'."
    return _git(command, path)


@tool("git_change", "Change a git repo: add, commit, stash, checkout, switch, restore, pull, fetch, init.",
      {"command": "e.g. 'add .' or 'commit -m \"message\"'", "path": "repo folder"}, risk="ask")
def git_change(command: str = "", path: str = ".") -> str:
    sub = (shlex.split(str(command).removeprefix("git ").strip()) or [""])[0]
    if sub not in GIT_WRITE:
        return f"git_change only runs {', '.join(sorted(GIT_WRITE))}."
    return _git(command, path)


# ----------------------------------------------------------------- web + apps
@tool("search_web", "Search the web and return the top links.", {"query": "what to search for"})
def search_web(query: str = "") -> str:
    from googlesearch import search
    links = list(search(str(query), num_results=6))
    return "\n".join(f"{i}. {u}" for i, u in enumerate(links, 1)) or "No results."


@tool("open_url", "Open a website in the default browser.", {"url": "the address"}, risk="ask")
def open_url(url: str = "") -> str:
    url = str(url).strip()
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    webbrowser.open(url)
    return f"Opened {url}"


@tool("open_app", "Open an app from Hybrid Agent's app list (see list_apps).",
      {"app": "the app key from list_apps"}, risk="ask")
def open_app(app: str = "", _user: str | None = None) -> str:
    import apps
    res = apps.launch(str(app).strip().lower(), _user)
    if res.get("ok"):
        return f"Opened {res.get('label', app)} (via {res.get('via')})."
    return res.get("error") or "Could not open that app."


@tool("list_apps", "List the apps Hybrid Agent can open.")
def list_apps(_user: str | None = None) -> str:
    import apps
    return "\n".join(f"{a.get('key')}: {a.get('label')}" for a in apps.catalogue(_user))


# ----------------------------------------------------------------- keyboard + WhatsApp
@tool("type_text", "Type text into whichever window is focused, after a 4 second pause so the user can click it.",
      {"text": "what to type", "press_enter": "true to press Enter after"}, risk="ask")
def type_text(text: str = "", press_enter=False) -> str:
    import pyautogui
    import pyperclip
    time.sleep(4)
    pyperclip.copy(str(text))
    pyautogui.hotkey("ctrl", "v")
    if str(press_enter).lower() in {"true", "1", "yes"}:
        time.sleep(0.1)
        pyautogui.press("enter")
    return f"Typed {len(str(text))} characters."


@tool("press_keys", "Press a keyboard shortcut such as 'ctrl+c', 'alt+tab' or 'win+d'.",
      {"keys": "keys joined by +"}, risk="ask")
def press_keys(keys: str = "") -> str:
    import pyautogui
    parts = [k.strip().lower() for k in str(keys).split("+") if k.strip()]
    if not parts:
        return "Say which keys to press."
    pyautogui.hotkey(*parts)
    return f"Pressed {' + '.join(parts)}"


def _whatsapp(contact: str):
    import pyautogui
    import pyperclip
    base = Path(os.environ.get("LOCALAPPDATA", ""))
    for exe in (base / "WhatsApp" / "WhatsApp.exe", base / "Programs" / "WhatsApp" / "WhatsApp.exe"):
        if exe.exists():
            os.startfile(str(exe))                       # noqa: S606 - fixed, known path
            break
    else:
        os.startfile("whatsapp:")                        # Microsoft Store build registers this protocol
    time.sleep(4)
    pyautogui.hotkey("ctrl", "f")
    time.sleep(0.5)
    pyperclip.copy(contact)
    pyautogui.hotkey("ctrl", "v")
    time.sleep(1.4)
    pyautogui.press("down")
    pyautogui.press("enter")
    time.sleep(1.0)
    return pyautogui, pyperclip


@tool("whatsapp_message", "Send a WhatsApp message from the desktop app.",
      {"contact": "contact name as saved", "message": "the text"}, risk="ask")
def whatsapp_message(contact: str = "", message: str = "") -> str:
    pyautogui, pyperclip = _whatsapp(str(contact))
    pyperclip.copy(str(message))
    pyautogui.hotkey("ctrl", "v")
    time.sleep(0.3)
    pyautogui.press("enter")
    return f"Sent the message to {contact}. Check WhatsApp to confirm it went to the right chat."


@tool("whatsapp_call", "Start a WhatsApp voice call from the desktop app.",
      {"contact": "contact name as saved"}, risk="ask")
def whatsapp_call(contact: str = "") -> str:
    pyautogui, _ = _whatsapp(str(contact))
    pyautogui.hotkey("ctrl", "shift", "c")
    return f"Calling {contact}."


# ----------------------------------------------------------------- vision
@tool("look_at_screen", "Take a screenshot and have Gemini describe it or answer a question about it.",
      {"question": "what to look for"}, risk="ask")
def look_at_screen(question: str = "Describe what is on my screen.", _user: str | None = None) -> str:
    from PIL import ImageGrab
    import providers

    key = providers.get_key("gemini", _user or providers.LOCAL_PROFILE)
    if not key:
        return "Looking at the screen needs a Gemini API key on this account (console > KEYS)."
    shot = ImageGrab.grab()
    if shot.width > 1600:
        shot = shot.resize((1600, int(shot.height * 1600 / shot.width)))
    buf = io.BytesIO()
    shot.convert("RGB").save(buf, "JPEG", quality=80)
    payload = {"contents": [{"role": "user", "parts": [
        {"inline_data": {"mime_type": "image/jpeg", "data": base64.b64encode(buf.getvalue()).decode()}},
        {"text": str(question)}]}]}
    req = urllib.request.Request(
        f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={key}",
        data=json.dumps(payload).encode(), headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=90) as r:
        data = json.loads(r.read().decode())
    parts = (data.get("candidates") or [{}])[0].get("content", {}).get("parts", [])
    return "".join(p.get("text", "") for p in parts).strip() or "Gemini returned nothing."


# ----------------------------------------------------------------- registry API
def available() -> dict[str, Tool]:
    return TOOLS if local_mode() else {}


def describe() -> list:
    return [{"name": t.name, "summary": t.summary, "params": t.params, "risk": t.risk}
            for t in available().values()]


def system_prompt() -> str:
    tools = available()
    lines = [
        "You are Nexus, the agent inside Hybrid Agent. You help the user by talking and by using tools "
        "on their Windows PC.",
        "",
    ]
    if tools:
        lines += [
            "To use a tool, reply with ONLY a fenced json block and nothing else:",
            "```json",
            '{"tool": "tool_name", "args": {"arg": "value"}}',
            "```",
            "You will get the result back as TOOL RESULT, then you can use another tool or answer.",
            "Use one tool at a time. Never invent tool results. If a tool fails, say so plainly.",
            "Tools marked (asks first) need the user's approval; if they deny it, accept that and move on.",
            "When you have the answer, reply in plain text (short, clear, no JSON).",
            "",
            "Tools:",
        ]
        for t in tools.values():
            args = ", ".join(f"{k}: {v}" for k, v in t.params.items()) or "no arguments"
            flag = " (asks first)" if t.risk == "ask" else ""
            lines.append(f"- {t.name}{flag}: {t.summary} Args: {args}")
        lines += ["", f"Files you create go to {OUT_DIR}. The Desktop is {_desktop()}."]
    else:
        lines.append("This copy runs in the cloud, so you have no access to the user's PC. "
                     "Just answer helpfully, and mention that PC actions need the local version.")
    return "\n".join(lines)


def parse_call(text: str) -> dict | None:
    """Pull a {"tool":..., "args":...} request out of a model reply, or None for a plain answer."""
    import re
    candidates = re.findall(r"```(?:json)?\s*(\{[\s\S]*?\})\s*```", text)
    if not candidates:
        stripped = text.strip()
        if stripped.startswith("{") and stripped.endswith("}"):
            candidates = [stripped]
    for c in candidates:
        try:
            obj = json.loads(c)
        except json.JSONDecodeError:
            continue
        if isinstance(obj, dict) and isinstance(obj.get("tool"), str):
            args = obj.get("args") if isinstance(obj.get("args"), dict) else {}
            return {"tool": obj["tool"], "args": args}
    return None


def transcript(messages: list) -> str:
    """Flatten the browser's agent conversation into one prompt for any provider."""
    out = []
    ctx = memory_context()
    if ctx:
        out.append(ctx)
    for m in messages[-30:]:
        role, text = m.get("role"), str(m.get("content", ""))[:MAX_RESULT]
        if role == "user":
            out.append(f"USER: {text}")
        elif role == "assistant":
            out.append(f"NEXUS: {text}")
        elif role == "tool":
            out.append(f"TOOL RESULT ({m.get('tool', '?')}): {text}")
    out.append("NEXUS:")
    return "\n\n".join(out)


def run(name: str, args: dict, user: str | None = None) -> dict:
    t = available().get(name)
    if not t:
        return {"ok": False, "result": f"There is no tool called '{name}'."}
    import inspect
    accepted = inspect.signature(t.fn).parameters
    kwargs = {k: v for k, v in (args or {}).items() if k in accepted and not k.startswith("_")}
    if "_user" in accepted:
        kwargs["_user"] = user
    try:
        out = str(t.fn(**kwargs))
        ok = True
    except Exception as exc:                            # noqa: BLE001 - reported to the model, not raised
        out, ok = f"{name} failed: {type(exc).__name__}: {exc}", False
    if len(out) > MAX_RESULT:
        out = out[:MAX_RESULT] + f"\n... ({len(out) - MAX_RESULT} more characters cut)"
    return {"ok": ok, "result": out}


def output_file(name: str) -> Path | None:
    """A file this agent wrote, by bare name only - never a path the caller chooses."""
    if not name or Path(name).name != name:
        return None
    p = OUT_DIR / name
    return p if p.is_file() else None
